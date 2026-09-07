from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageEnhance

OUT = Path('teslim')
OUT.mkdir(exist_ok=True)
LOGO = Path(r'C:\Users\eyupm\Downloads\1630586807885.jpg')
LOGO_BW = Path('logo_siyah_beyaz.png')
TODAY = '21 Temmuz 2026'
VESSELS = ['SNP GALAXY', 'SNP SKY', 'SNP SPACE', 'SNP STAR', 'SNP WIND', 'SNP PAZAR']
# Siyah-beyaz baskı için yüksek kontrastlı nötr palet.
NAVY = '222222'; BLUE = '3F3F3F'; CYAN = 'E2E2E2'; PALE = 'F2F2F2'; GRAY = '666666'; RED = '111111'; AMBER = '555555'; GREEN = '333333'; WHITE = 'FFFFFF'; INK='111111'
NOTICE = "Bu belgeler stajyer öğrencimiz Eyüp Miraç OY'un operasyonları daha iyi öğrenebilmesi ve öğrendiklerini uygulayabilmesi  adına şirket tarafından özel izinle hazırlanmıştır. Üçüncü kişilerce kullanılması ya da paylaşılması yasaktır."

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(cell, color='C8C8C8', size='4'):
    tcPr = cell._tc.get_or_add_tcPr(); b = tcPr.first_child_found_in('w:tcBorders')
    if b is None: b=OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),size); e.set(qn('w:color'),color); b.append(e)

def margins(cell, top=90, start=110, bottom=90, end=110):
    tcPr=cell._tc.get_or_add_tcPr(); m=tcPr.first_child_found_in('w:tcMar')
    if m is None: m=OxmlElement('w:tcMar'); tcPr.append(m)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        x=OxmlElement('w:'+tag); x.set(qn('w:w'),str(val)); x.set(qn('w:type'),'dxa'); m.append(x)

def set_repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); x=OxmlElement('w:tblHeader'); x.set(qn('w:val'),'true'); trPr.append(x)

def cant_split(row):
    trPr=row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))

def set_cell_text(cell, text, bold=False, color=INK, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text=''; p=cell.paragraphs[0]; p.alignment=align; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
    r=p.add_run(str(text)); r.bold=bold; r.font.name='Arial'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell); borders(cell)

def set_width(cell, dxa):
    tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.find(qn('w:tcW'))
    if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'),str(dxa)); tcW.set(qn('w:type'),'dxa')

def table(doc, headers, rows, widths, header_fill=NAVY, font=8.2):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    grid=t._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'),str(w)); grid.append(gc)
    tblPr=t._tbl.tblPr; tblW=tblPr.find(qn('w:tblW')); tblW.set(qn('w:w'),str(sum(widths))); tblW.set(qn('w:type'),'dxa')
    ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),'110'); ind.set(qn('w:type'),'dxa'); tblPr.append(ind)
    for i,h in enumerate(headers): set_width(t.rows[0].cells[i], widths[i]); set_cell_text(t.rows[0].cells[i],h,True,WHITE,font); shade(t.rows[0].cells[i],header_fill)
    set_repeat_header(t.rows[0]); cant_split(t.rows[0])
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells; cant_split(t.rows[-1])
        for i,v in enumerate(row):
            set_width(cells[i],widths[i]); set_cell_text(cells[i],v,False,INK,font)
            if ridx%2: shade(cells[i],PALE)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def add_page_field(p):
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=p.add_run('Sayfa '); r.font.name='Arial'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)

def setup(kind, code):
    d=Document(); s=d.sections[0]
    s.page_width=Inches(8.5); s.page_height=Inches(11); s.top_margin=Inches(.7); s.bottom_margin=Inches(.85); s.left_margin=Inches(.75); s.right_margin=Inches(.75); s.header_distance=Inches(.25); s.footer_distance=Inches(.28)
    styles=d.styles
    normal=styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(9.5); normal.font.color.rgb=RGBColor.from_string(INK); normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.08
    for name,size,before,after in [('Heading 1',16,14,7),('Heading 2',12.5,10,5),('Heading 3',10.5,7,3)]:
        st=styles[name]; st.font.name='Arial'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(NAVY); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
    h=s.header; p=h.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    r=p.add_run(f'SINOP SHIPPING CORP.  |  {kind}'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(NAVY)
    f=s.footer; p=f.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
    r=p.add_run(NOTICE); r.font.name='Arial'; r.font.size=Pt(6.5); r.font.color.rgb=RGBColor.from_string(GRAY); r.italic=True
    add_page_field(f.add_paragraph())
    d.core_properties.title=kind; d.core_properties.subject='Şirket içi eğitim amaçlı ön çalışma'; d.core_properties.author='Sinop Shipping Corp. - Eğitim Amaçlı Çalışma'
    return d

def title_page(d, title, subtitle, code):
    logo_path = LOGO_BW if LOGO_BW.exists() else LOGO
    if logo_path.exists():
        p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(22); p.paragraph_format.space_after=Pt(12); p.add_run().add_picture(str(logo_path), width=Inches(1.55))
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(5)
    r=p.add_run(title); r.bold=True; r.font.name='Arial'; r.font.size=Pt(24); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(subtitle); r.font.name='Arial'; r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(BLUE)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(20)
    r=p.add_run('ŞİRKET İÇİ • EĞİTİM AMAÇLI • KONTROLLÜ KOPYA • SİYAH-BEYAZ BASKI'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(RED)
    table(d,['BELGE BİLGİSİ','DEĞER'],[
        ('Şirket','Sinop Shipping Corp.'),('Rapor kapsamı','Mevcut filo: 6 gemi'),('Rapor tarihi',TODAY),('Belge kodu',code),('Revizyon','00'),('Hazırlama amacı','Stajyer operasyon eğitimi / şirket içi çalışma')
    ],[2200,7160],font=9)
    d.add_paragraph('Belge Statüsü',style='Heading 2')
    p=d.add_paragraph('Bu çalışma, yayımlanmış filo ve iletişim bilgileri ile genel denizcilik risk/KPI uygulamalarına dayanır. Operasyonel gerçekleşen değerler, gemi sertifikaları ve şirket içi kayıtlar paylaşılmadığından belge bir denetim sonucu, klas raporu veya yetkili şirket beyanı değildir. Şirket içi verilerle doğrulanarak kullanılmalıdır.')
    p.paragraph_format.space_after=Pt(8)
    d.add_page_break()

def contact_block(d):
    d.add_heading('Kurumsal İletişim ve Belge Dağıtımı',level=1)
    table(d,['BİRİM','İLETİŞİM'],[
        ('Merkez adres','Postane Mah., İdil Sk. No:9 B/1, 34940 Tuzla - İstanbul / Türkiye'),
        ('Telefon','+90 216 504 83 55'),('Chartering','sinop@sinopshipping.com'),('Operation','operation@sinopshipping.com'),('Technical','info@sinopshipping.com'),('Crew','crew@sinopshipping.com'),('Web','www.sinopshipping.com'),('Çalışma saatleri','Pazartesi-Cuma, 09:00-18:00')
    ],[2200,7160],font=8.8)
    p=d.add_paragraph('Kaynak notu: Kurumsal adres, iletişim ve filo adları Sinop Shipping Corp. resmî internet sitesinden 21.07.2026 tarihinde doğrulanmıştır.'); p.style='Caption'

def risk_report():
    d=setup('Filo Risk Analizi Raporu','SSC-RA-2026-01')
    title_page(d,'FİLO RİSK ANALİZİ RAPORU','SNP GALAXY • SNP SKY • SNP SPACE • SNP STAR • SNP WIND • SNP PAZAR','SSC-RA-2026-01')
    d.add_heading('1. Yönetici Özeti',level=1)
    d.add_paragraph('Bu ön çalışma, altı gemi için ortak operasyonel tehlikeleri tanımlar ve gemi bazında takip edilebilir bir risk kayıt yapısı kurar. Paylaşılan olay, bakım, sefer, klas veya PSC verisi bulunmadığından risk puanları “başlangıç maruziyet puanı”dır; gerçek performans göstergesi değildir.')
    table(d,['ALAN','ÖN DEĞERLENDİRME','ÖNCELİK'],[
        ('Seyir ve çatışma','Yoğun trafik, kısıtlı görüş, köprüüstü kaynak yönetimi','Yüksek'),('Yük operasyonları','İstif, bağlama, ambar emniyeti, ağır/proje yükü','Yüksek'),('Makine ve enerji','Ana makine/yardımcı sistem arızası, plansız duruş','Yüksek'),('İSG','Kapalı mahal, yüksekte çalışma, sıcak çalışma','Yüksek'),('Çevre','Yakıt/yağ döküntüsü, atık, emisyon uyumu','Orta-Yüksek'),('Siber/güvenlik','GNSS/AIS etkilenmesi, erişim ve veri güvenliği','Orta')
    ],[2100,5360,1900],font=8.3)
    d.add_heading('2. Yöntem ve Ölçek',level=1)
    d.add_paragraph('Risk puanı = Olasılık (1-5) × Etki (1-5). Kalan risk, mevcut/önerilen kontroller uygulandıktan sonra yeniden puanlanır. Her gemi için Kaptan ve Başmühendis, gerçek kayıtları kullanarak puanları doğrulamalıdır.')
    table(d,['PUAN','SEVİYE','YÖNETİM KURALI'],[('1-4','Düşük','Rutin kontrol ve izleme'),('5-9','Orta','Planlı iyileştirme ve sorumlu atama'),('10-16','Yüksek','Öncelikli kontrol; yönetim takibi'),('17-25','Kritik','Faaliyet öncesi risk düşürme / durdurma değerlendirmesi')],[1500,1900,5960],font=8.5)
    d.add_heading('3. Kapsam ve Varsayımlar',level=1)
    d.add_paragraph('Kapsam; seyir, yük, makine, personel, çevre, güvenlik, siber risk ve mevzuat uyumudur. Gemi tipi/yaşı/klası/tonajı, son seferler ve bakım geçmişi doğrulanmadığı için gemiler hakkında özgül teknik iddia kurulmamıştır. Aynı başlangıç riskleri, gemi bazında veri girilmesi için kontrol tabanı olarak kullanılmıştır.')
    d.add_page_break()
    base=[
      ('Seyir / çatışma','Yoğun trafik, kısıtlı görüş veya yorgunluk','Çatışma, karaya oturma, can/çevre zararı','3','5','15 Yüksek','Geçiş planı, BRM, ECDIS/radar çapraz kontrolü','2×5=10'),
      ('Yük operasyonu','Eksik istif/bağlama planı, uygunsuz ekipman','Yük kayması, hasar, yaralanma','3','5','15 Yüksek','Onaylı plan, toolbox, ekipman SWL ve saha gözetimi','2×5=10'),
      ('Makine arızası','PMS gecikmesi, kritik yedek eksikliği','Tahrik kaybı, gecikme, kurtarma ihtiyacı','3','5','15 Yüksek','PMS, kritik ekipman testi, yedek ve trend izleme','2×5=10'),
      ('İş sağlığı ve güvenliği','Kapalı mahal/sıcak iş/yüksekte çalışma','Ağır yaralanma veya ölüm','3','5','15 Yüksek','PTW, gaz ölçümü, LOTO, kurtarma planı','1×5=5'),
      ('Yangın','Yakıt/sıcak yüzey/elektrik kaynağı','Can kaybı, gemi ve yük hasarı','2','5','10 Yüksek','Yangın devriyesi, dedektör testleri, tatbikat','1×5=5'),
      ('Çevre kirliliği','Bunker/transfer/atık operasyonu','Döküntü, ceza, itibar kaybı','3','4','12 Yüksek','SOPEP, scupper kapama, transfer checklisti','1×4=4'),
      ('Siber ve GNSS','Yetkisiz erişim, sahte konum/sinyal kesintisi','Seyir/veri bütünlüğü kaybı','2','4','8 Orta','MFA, yedekleme, görsel/radar doğrulama, olay planı','1×4=4'),
      ('Mevzuat / PSC','Sertifika veya kayıt uygunsuzluğu','Tutuklama, gecikme, ticari kayıp','2','5','10 Yüksek','Pre-arrival kontrolü, sertifika matrisi, iç denetim','1×5=5')]
    for idx,v in enumerate(VESSELS,1):
        d.add_heading(f'4.{idx} {v} - Risk Kayıt Kartı',level=1)
        d.add_paragraph('Durum: Ön değerlendirme / gemi kayıtlarıyla doğrulama bekleniyor. Sorumlu önerisi: Kaptan, Başmühendis ve ilgili şirket operasyon/teknik sorumlusu.')
        table(d,['RİSK','TEHLİKE / NEDEN','OLASI SONUÇ','O','E','İLK PUAN','ANA KONTROLLER','KALAN'],base,[1100,1580,1580,330,330,780,2700,960],font=6.8)
        d.add_paragraph('Gemiye özgü doğrulama alanları: son PSC/klas bulguları, 12 aylık arıza ve olay kayıtları, çalışma/dinlenme ihlalleri, kritik yedek durumu, son tatbikat sonuçları, rota ve yük profili.',style='Caption')
        if idx != len(VESSELS): d.add_page_break()
    d.add_page_break(); d.add_heading('5. Filo Düzeyi Öncelikli Aksiyon Planı',level=1)
    table(d,['NO','AKSİYON','SORUMLU','TERMİN','KANIT / KAPANIŞ'],[
      ('RA-01','Her geminin son 12 aylık olay, ramak kala ve arıza verisini konsolide et','Operation / Technical','30 gün','Onaylı veri özeti'),('RA-02','Kritik ekipman ve yedek parça listesini gemi bazında doğrula','Technical / C/E','30 gün','PMS ve stok çıktısı'),('RA-03','Seyir ve yük operasyonu için hedefli toolbox / vaka çalışması yap','Master / Operation','Her sefer','Toplantı kaydı'),('RA-04','Kapalı mahal, sıcak iş ve LOTO izinlerini örneklemle denetle','DPA / Master','Aylık','İç denetim formu'),('RA-05','Pre-arrival PSC kontrolü ve sertifika matrisini standardize et','DPA / Master','Her liman','Tamamlanmış checklist'),('RA-06','Risk kayıtlarını olay ve KPI trendlerine göre üç ayda bir güncelle','Management','Çeyreklik','Revize risk kaydı')
    ],[600,3900,1550,1100,2210],font=8)
    d.add_heading('6. Onay ve Revizyon Kaydı',level=1)
    table(d,['ROL','AD SOYAD','TARİH','İMZA'],[('Hazırlayan / İnceleyen','','',''),('Operasyon Onayı','','',''),('Teknik Onay','','',''),('Yönetim Onayı','','','')],[2200,2900,1600,2660],font=8.5)
    contact_block(d)
    d.save(OUT/'Sinop_Shipping_Filo_Risk_Analizi_Siyah_Beyaz.docx')

def kpi_report():
    d=setup('Filo KPI Performans Raporu','SSC-KPI-2026-01')
    title_page(d,'FİLO KPI PERFORMANS RAPORU','Hedefler • Hesaplama Yöntemleri • Gemi Bazlı Skor Kartları','SSC-KPI-2026-01')
    d.add_heading('1. Yönetici Özeti',level=1)
    d.add_paragraph('Bu rapor, Sinop Shipping filosu için ölçülebilir performans yönetimi çerçevesi sunar. Gerçek operasyonel veri sağlanmadığından herhangi bir gemi için gerçekleşen başarı oranı üretilmemiştir. Tüm sonuç hücreleri “Veri bekleniyor” olarak işaretlenmiş ve hesaplama formülleri tanımlanmıştır.')
    table(d,['BOYUT','AĞIRLIK','AMAÇ'],[('Emniyet ve İSG','25%','Olayları ve yüksek potansiyelli sapmaları azaltmak'),('Teknik güvenilirlik','25%','Plansız duruşu ve kritik arızayı azaltmak'),('Operasyon ve ticari','20%','Sefer ve liman operasyonlarını planlı yürütmek'),('Çevre ve enerji','15%','Tüketim/emisyon/atık performansını izlemek'),('Uyum ve insan','15%','PSC, eğitim, yorgunluk ve denetim disiplinini güçlendirmek')],[3000,1200,5160],font=8.5)
    d.add_heading('2. Raporlama İlkeleri',level=1)
    d.add_paragraph('Aylık gemi raporu, şirket operasyon ve teknik kayıtlarıyla eşleştirilir. Payda sıfırsa KPI “uygulanamaz” kaydedilir. Düzeltme yapılan veriler revizyon iziyle saklanır. Sonuçlar hedefe göre Yeşil, Sarı veya Kırmızı sınıflandırılır; “Veri bekleniyor” ayrı statüdür ve performans sonucu sayılmaz.')
    d.add_heading('3. Filo KPI Sözlüğü ve Hedefleri',level=1)
    rows=[
      ('KPI-01','LTIF','(Kayıp günlü yaralanma × 1.000.000) / çalışma saati','0','Aylık'),('KPI-02','TRCF','(Kayıtlı vaka × 1.000.000) / çalışma saati','0','Aylık'),('KPI-03','Ramak kala bildirimi','Adet / gemi-ay','≥2','Aylık'),('KPI-04','Kritik arıza','Kritik ekipman arıza adedi','0','Aylık'),('KPI-05','Teknik kullanılabilirlik','Çalışabilir saat / planlı saat ×100','≥98%','Aylık'),('KPI-06','PMS zamanında tamamlama','Zamanında kapanan iş / vadesi gelen iş ×100','≥95%','Aylık'),('KPI-07','Sefer zamanında performans','Plan dahilinde tamamlanan sefer / toplam ×100','≥90%','Aylık'),('KPI-08','Liman kalış sapması','(Gerçek-plan) saat / plan saat ×100','≤10%','Sefer'),('KPI-09','Yakıt yoğunluğu','Toplam yakıt (t) / deniz mili veya ton-mil','Baz çizgiye göre ↓','Sefer'),('KPI-10','CO₂ yoğunluğu','CO₂ (t) / ton-mil','Baz çizgiye göre ↓','Aylık'),('KPI-11','PSC deficiency','Bulgu adedi / denetim','0 majör; toplam ↓','Denetim'),('KPI-12','PSC detention','Tutuklama adedi','0','Denetim'),('KPI-13','Eğitim/tatbikat tamamlama','Zamanında tamamlanan / planlanan ×100','100%','Aylık'),('KPI-14','Çalışma-dinlenme uygunsuzluğu','İhlal adedi','0','Aylık'),('KPI-15','Müşteri/yük hasar talebi','Haklı talep adedi','0','Sefer')]
    table(d,['KOD','GÖSTERGE','FORMÜL / TANIM','HEDEF','SIKLIK'],rows,[800,1800,3780,1700,1280],font=7.4)
    d.add_page_break(); d.add_heading('4. Filo Konsolide Gösterge Paneli',level=1)
    fleetrows=[]
    for v in VESSELS: fleetrows.append((v,'Veri bekleniyor','Veri bekleniyor','Veri bekleniyor','Veri bekleniyor','Veri bekleniyor'))
    table(d,['GEMİ','EMNİYET','TEKNİK','OPERASYON','ÇEVRE','UYUM/İNSAN'],fleetrows,[1760,1520,1520,1520,1520,1520],font=7.8)
    d.add_paragraph('Not: Renkli skor yalnızca kaynak kayıtları doğrulandıktan sonra hesaplanır. Veri yokluğu “başarılı” kabul edilmez.',style='Caption')
    for idx,v in enumerate(VESSELS,1):
        d.add_page_break(); d.add_heading(f'5.{idx} {v} - Aylık KPI Skor Kartı',level=1)
        d.add_paragraph('Raporlama dönemi: ____ / 2026    Veri sorumlusu: ____    Gemi onayı: ____    Şirket onayı: ____')
        score=[]
        selected=[rows[i] for i in [0,1,2,3,4,5,6,7,8,9,10,11,12,13,14]]
        for code,name,formula,target,freq in selected: score.append((code,name,target,'Veri bekleniyor','—','Gemi/şirket kayıtları'))
        table(d,['KOD','KPI','HEDEF','GERÇEKLEŞEN','DURUM','KANIT KAYNAĞI'],score,[720,1900,1400,1700,900,2740],font=7.3)
        d.add_heading('Yönetim Yorumu ve Düzeltici Faaliyet',level=2)
        table(d,['SAPMA / BULGU','KÖK NEDEN','AKSİYON','SORUMLU','TERMİN'],[('','','','',''),('','','','',''),('','','','','')],[1900,1900,2500,1500,1560],font=8)
    d.add_page_break(); d.add_heading('6. Skorlama ve Trafik Işığı Kuralı',level=1)
    table(d,['STATÜ','KURAL','YÖNETİM BEKLENTİSİ'],[('Yeşil','Hedef karşılandı','Standart sürdürülür; iyi uygulama paylaşılır'),('Sarı','Hedeften sınırlı sapma / kötüleşen trend','Kök neden ve tarihli aksiyon'),('Kırmızı','Kritik eşik ihlali veya tekrarlayan sapma','Yönetim eskalasyonu ve yakın takip'),('Gri','Eksik veya doğrulanmamış veri','Veri tamamlanmadan performans hükmü verilmez')],[1400,3100,4860],font=8.5)
    d.add_heading('7. Veri Kaynakları ve Kontrol',level=1)
    table(d,['KPI GRUBU','BİRİNCİL KAYNAK','KONTROL'],[('Emniyet/İSG','Olay, ramak kala, çalışma saati kayıtları','DPA / HSEQ doğrulaması'),('Teknik','PMS, defect list, engine log, spare list','Technical Superintendent'),('Operasyon','Noon report, SOF, sefer planı, liman kayıtları','Operation Department'),('Çevre/enerji','Bunker ROB, BDN, mesafe/yük ve atık kayıtları','Master / C/E / Office'),('Uyum/insan','PSC/flag/class, eğitim ve work-rest kayıtları','DPA / Crew / Master')],[2200,3860,3300],font=8.2)
    d.add_heading('8. Onay ve Revizyon Kaydı',level=1)
    table(d,['ROL','AD SOYAD','TARİH','İMZA'],[('Hazırlayan / İnceleyen','','',''),('Operasyon Onayı','','',''),('Teknik Onay','','',''),('Yönetim Onayı','','','')],[2200,2900,1600,2660],font=8.5)
    contact_block(d)
    d.save(OUT/'Sinop_Shipping_Filo_KPI_Raporu_Siyah_Beyaz.docx')

if __name__=='__main__':
    if LOGO.exists():
        with Image.open(LOGO) as im:
            gray = ImageEnhance.Contrast(im.convert('L')).enhance(1.25)
            gray.save(LOGO_BW)
    risk_report(); kpi_report(); print('created', *OUT.glob('*.docx'), sep='\n')
