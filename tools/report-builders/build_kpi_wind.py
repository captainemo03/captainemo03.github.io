from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_reports import setup, table, OUT, LOGO_BW, LOGO, NAVY, GRAY

def build():
    d=setup('Aylık KPI Performans Kartı - SNP WIND','SSC-KPI-WND-2026-08')
    s=d.sections[0]; s.top_margin=Inches(.40); s.bottom_margin=Inches(.73); s.left_margin=Inches(.58); s.right_margin=Inches(.58)
    d.styles['Normal'].font.size=Pt(7.8); d.styles['Normal'].paragraph_format.space_after=Pt(1.5)
    logo=LOGO_BW if LOGO_BW.exists() else LOGO
    if logo.exists():
        p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); p.add_run().add_picture(str(logo),width=Inches(.68))
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
    r=p.add_run('AYLIK KPI PERFORMANS KARTI | SNP WIND'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(15.5); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(3)
    r=p.add_run('KURGUSAL EĞİTİM SENARYOSU • AĞUSTOS 2026 • SİYAH-BEYAZ BASKI'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(7.1); r.font.color.rgb=RGBColor.from_string(GRAY)
    table(d,['RAPOR BİLGİSİ','DEĞER','RAPOR BİLGİSİ','DEĞER'],[
      ('Gemi','SNP WIND','Rapor dönemi','01–31 Ağustos 2026'),('Ana ticaret hattı','Cenova–İzmir / genel kargo','Tamamlanan sefer','5 sefer'),('Toplam çalışma saati','4.800 kişi-saat','Toplam seyir mesafesi','6.420 NM'),('Belge / Revizyon','SSC-KPI-WND-2026-08 / 00','Genel skor','78 / 100 — SARI')],[1750,2930,1750,2930],font=7.0)
    p=d.add_paragraph('KPI SONUÇLARI'); p.style='Heading 2'; p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(1)
    table(d,['KOD','GÖSTERGE','HESAP / VERİ','HEDEF','GERÇEKLEŞEN','DURUM'],[
      ('KPI-01','LTIF','0 LTI × 1.000.000 / 4.800 saat','0','0,00','YEŞİL'),
      ('KPI-02','TRCF','1 kayıtlı vaka × 1.000.000 / 4.800 saat','0','208,33','KIRMIZI'),
      ('KPI-03','Ramak kala bildirimi','Aylık bildirim adedi','≥2','3 adet','YEŞİL'),
      ('KPI-04','Kritik ekipman arızası','Kritik arıza adedi','0','0 adet','YEŞİL'),
      ('KPI-05','Teknik kullanılabilirlik','710 çalışabilir / 720 planlı saat','≥98%','98,61%','YEŞİL'),
      ('KPI-06','PMS zamanında tamamlama','81 zamanında / 86 vadesi gelen iş','≥95%','94,19%','SARI'),
      ('KPI-07','Sefer zamanında performans','4 zamanında / 5 toplam sefer','≥90%','80,00%','KIRMIZI'),
      ('KPI-08','Liman kalış sapması','Gerçek-plan farkı / planlı süre','≤10%','8,00%','YEŞİL'),
      ('KPI-09','Yakıt yoğunluğu','1.130 t / 6.420 NM','≤0,185 t/NM','0,176 t/NM','YEŞİL'),
      ('KPI-10','CO₂ yoğunluğu','Sefer/yük/mesafe verisi','≤12,8 g/t-NM','12,4 g/t-NM','YEŞİL'),
      ('KPI-11','PSC bulgusu','1 denetimde 2 minör bulgu','0','2 minör','SARI'),
      ('KPI-12','PSC detention','Tutuklama adedi','0','0','YEŞİL'),
      ('KPI-13','Eğitim/tatbikat tamamlama','11 tamamlanan / 12 planlanan','100%','91,67%','SARI'),
      ('KPI-14','Çalışma-dinlenme uygunsuzluğu','Aylık doğrulanmış ihlal','0','2 ihlal','KIRMIZI'),
      ('KPI-15','Yük hasar talebi','Haklı müşteri talebi','0','0','YEŞİL')],[650,1800,2900,1350,1500,1160],font=6.1)
    p=d.add_paragraph('BOYUT BAZLI SKOR'); p.style='Heading 2'; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(1)
    table(d,['BOYUT','AĞIRLIK','SKOR','AĞIRLIKLI PUAN','YORUM'],[
      ('Emniyet ve İSG','25%','68/100','17,00','1 kayıtlı vaka; ramak kala hedefi karşılandı'),('Teknik güvenilirlik','25%','91/100','22,75','Kritik arıza yok; PMS sınıra yakın'),('Operasyon / ticari','20%','74/100','14,80','Zamanında sefer oranı hedef altında'),('Çevre / enerji','15%','92/100','13,80','Yakıt ve CO₂ yoğunluğu baz çizginin altında'),('Uyum / insan','15%','64/100','9,60','PSC minör bulguları ve work-rest ihlalleri')],[2100,1050,1050,1650,3510],font=6.4)
    p=d.add_paragraph('TOPLAM AĞIRLIKLI SKOR: 77,95 ≈ 78 / 100 — SARI (İYİLEŞTİRME GEREKLİ)'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].bold=True; p.paragraph_format.space_after=Pt(2)
    p=d.add_paragraph('DÜZELTİCİ FAALİYETLER'); p.style='Heading 2'; p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(1)
    table(d,['BULGU / SAPMA','AKSİYON','SORUMLU','TERMİN'],[
      ('TRCF: 1 kayıtlı vaka','Olay kök neden analizi; ilgili işe toolbox ve saha gözlemi','Master / DPA','7 gün'),('PMS: %94,19','Geciken 5 işi kritikliye göre planla ve kapat','Chief Engineer / Technical','15 gün'),('Sefer OTP: %80','Gecikme nedenlerini liman, hava ve teknik olarak ayrıştır','Operation','Aylık kapanış'),('2 work-rest ihlali','Vardiya planını revize et; haftalık uygunluk kontrolü','Master / Crew','Derhal')],[2300,4200,1700,1160],font=6.3)
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(1)
    r=p.add_run('Veri Notu: '); r.bold=True
    p.add_run('Tüm değerler staj eğitimi için oluşturulmuş kurgusal örnektir; gerçek SNP WIND performansını veya şirket kaydını temsil etmez. Gerçek kullanımda noon report, PMS, PSC, olay ve work-rest kayıtlarıyla doğrulanmalıdır.')
    table(d,['HAZIRLAYAN','MASTER / GEMİ ONAYI','OPERATION / TECHNICAL ONAYI'],[('Eyüp Miraç OY\nTarih / İmza: __________','Ad Soyad / İmza:\n________________','Ad Soyad / İmza:\n________________')],[3000,3000,3360],font=6.6)
    path=OUT/'SNP_WIND_KPI_Raporu_Doldurulmus_Ornek.docx'; d.save(path); print(path)

if __name__=='__main__': build()
