from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_reports import setup, table, OUT, LOGO_BW, LOGO, NAVY, GRAY

def build():
    d = setup('ETA-ETD Hesaplama Formu - SNP GALAXY', 'SSC-OPS-ETA-01-GXY')
    s=d.sections[0]; s.top_margin=Inches(.45); s.bottom_margin=Inches(.75); s.left_margin=Inches(.65); s.right_margin=Inches(.65)
    d.styles['Normal'].font.size=Pt(8.3); d.styles['Normal'].paragraph_format.space_after=Pt(1.5)
    logo=LOGO_BW if LOGO_BW.exists() else LOGO
    if logo.exists():
        p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); p.add_run().add_picture(str(logo),width=Inches(.78))
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
    r=p.add_run('ETA - ETD HESAPLAMA FORMU | SNP GALAXY'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(16); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(3)
    r=p.add_run('KURGUSAL EGITIM SENARYOSU - SIYAH-BEYAZ BASKI'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(7.4); r.font.color.rgb=RGBColor.from_string(GRAY)
    table(d,['BELGE / SEFER BILGISI','DEGER','BELGE / SEFER BILGISI','DEGER'],[
      ('Belge Kodu','SSC-OPS-ETA-01-GXY','Tarih / Revizyon','21.07.2026 | 00'),('Gemi','SNP GALAXY','Sefer No.','GXY-TRG-2026/07'),('Kalkis Limani','Constanta, Romanya','Varis Limani','Pire, Yunanistan'),('Rota','Karadeniz-Bogazlar-Ege','Saat Dilimi','UTC')],[1900,2780,1900,2780],font=7.5)
    p=d.add_paragraph('1. ETA HESABI - TAHMINI VARIS ZAMANI'); p.style='Heading 2'; p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(1)
    table(d,['GIRDI','DEGER','BIRIM / ACIKLAMA'],[
      ('Kalkis tarihi ve saati','22.07.2026 - 08:00','UTC'),('Toplam rota mesafesi (D)','530','Deniz mili - NM'),('Ortalama seyir surati (V)','10,5','Knot - NM/saat'),('Hava/deniz gecikmesi','4,0','Saat'),('Trafik / Bogaz gecis gecikmesi','5,0','Saat'),('Diger emniyet payi','1,5','Saat')],[3200,2800,3360],font=7.5)
    table(d,['HESAPLAMA ADIMI','FORMUL','SONUC'],[
      ('Saf seyir suresi','530 NM / 10,5 kn','50 saat 29 dakika'),('Toplam gecikme','4,0 + 5,0 + 1,5','10 saat 30 dakika'),('Duzeltilmis seyir suresi','50 sa 29 dk + 10 sa 30 dk','60 saat 59 dakika'),('ETA','22.07.2026 08:00 + 60 sa 59 dk','24.07.2026 - 20:59 UTC')],[2800,3900,2660],font=7.4)
    p=d.add_paragraph('2. ETD HESABI - TAHMINI KALKIS ZAMANI'); p.style='Heading 2'; p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(1)
    table(d,['LIMAN OPERASYONU','TAHMINI SURE','ACIKLAMA'],[
      ('Varis / yanasma beklemesi','6,0 saat','Rihtim sirasi ve liman trafigi'),('Yanasma ve hazirlik','2,0 saat','Pilotaj, palamar ve emniyet kontrolu'),('Tahliye operasyonu','30,0 saat','Planlanan genel kargo tahliyesi'),('Survey / evrak / clearance','3,0 saat','Draft survey ve liman evraklari'),('Kalkis hazirligi ve emniyet payi','1,5 saat','Pilot ve makine hazirligi')],[3300,2500,3560],font=7.2)
    table(d,['HESAPLAMA ADIMI','FORMUL','SONUC'],[
      ('Toplam liman suresi','6 + 2 + 30 + 3 + 1,5','42 saat 30 dakika'),('ETD','24.07.2026 20:59 + 42 sa 30 dk','26.07.2026 - 15:29 UTC')],[2800,3900,2660],font=7.4)
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(1); r=p.add_run('Kontrol Notu: '); r.bold=True
    p.add_run('Bu degerler yalnizca staj egitimi icin olusturulmus kurgusal bir senaryodur; gercek SNP GALAXY seferi veya sirket operasyon kaydi degildir. Rota, surat, hava, Bogaz trafigi ve liman programi degistiginde ETA/ETD yeniden hesaplanmalidir.')
    table(d,['HAZIRLAYAN','MASTER / GEMI ONAYI','OPERATION ONAYI'],[('Eyup Mirac OY\nTarih / Imza: __________','Ad Soyad / Imza:\n________________','Ad Soyad / Imza:\n________________')],[3120,3120,3120],font=7.2)
    path=OUT/'SNP_GALAXY_ETA_ETD_Hesaplama_Doldurulmus_Ornek.docx'; d.save(path); print(path)

if __name__=='__main__': build()
