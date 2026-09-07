from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_reports import setup, table, OUT, LOGO_BW, LOGO, NAVY, GRAY


def build():
    d = setup('Demurrage ve Dispatch Hesaplama Formu', 'SSC-OPS-DD-01')
    s = d.sections[0]
    s.top_margin = Inches(0.46); s.bottom_margin = Inches(0.75)
    s.left_margin = Inches(0.65); s.right_margin = Inches(0.65)
    d.styles['Normal'].font.size = Pt(8.3)
    d.styles['Normal'].paragraph_format.space_after = Pt(2)

    logo = LOGO_BW if LOGO_BW.exists() else LOGO
    if logo.exists():
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(logo), width=Inches(0.76))
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run('DEMURRAGE VE DISPATCH / DESPATCH HESAPLAMA FORMU')
    r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(14.5)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('ŞİRKET İÇİ • EĞİTİM AMAÇLI • SİYAH-BEYAZ BASKI')
    r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(7.2)
    r.font.color.rgb = RGBColor.from_string(GRAY)

    table(d, ['SEFER / SÖZLEŞME', 'DEĞER', 'SEFER / SÖZLEŞME', 'DEĞER'], [
        ('Gemi', '☐ GALAXY ☐ SKY ☐ SPACE', 'Devam', '☐ STAR ☐ WIND ☐ PAZAR'),
        ('Sefer / C/P tarihi', '', 'Liman / Terminal', ''),
        ('Yük / Miktar', '________________ MT', 'Operasyon', '☐ Yükleme  ☐ Tahliye'),
        ('Laytime şartı', '☐ SHINC ☐ SHEX ☐ Diğer: ___', 'Reversible', '☐ Evet  ☐ Hayır'),
    ], [1750, 2930, 1750, 2930], font=7.2)

    p = d.add_paragraph('1. LAYTIME SONUÇ GİRDİLERİ'); p.style = 'Heading 2'
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(1)
    table(d, ['HESAP KALEMİ', 'GÜN', 'SAAT', 'TOPLAM SAAT', 'AÇIKLAMA'], [
        ('İzin verilen laytime (A)', '', '', '', 'C/P yük miktarı ve oranından'),
        ('Kullanılan net laytime (B)', '', '', '', 'Time sheet sonrası sayılan süre'),
        ('Zaman farkı (B − A)', '', '', '', '(+) Demurrage / (−) Dispatch'),
    ], [2650, 1000, 1000, 1600, 3110], font=7.4)

    p = d.add_paragraph('2. SÖZLEŞMESEL ORANLAR'); p.style = 'Heading 2'
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(1)
    table(d, ['ORAN', 'GÜNLÜK TUTAR', 'SAATLİK TUTAR', 'PARA BİRİMİ / NOT'], [
        ('Demurrage', '____________ / gün', 'Günlük oran ÷ 24 = ____________', '☐ USD ☐ EUR ☐ Diğer: ____'),
        ('Dispatch / Despatch', '____________ / gün', 'Günlük oran ÷ 24 = ____________', '☐ USD ☐ EUR ☐ Diğer: ____'),
    ], [2200, 2100, 2800, 2260], font=7.5)

    p = d.add_paragraph('3. NİHAİ TUTAR HESABI'); p.style = 'Heading 2'
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(1)
    table(d, ['SONUÇ TÜRÜ', 'UYGULANACAK FORMÜL', 'HESAP', 'NİHAİ TUTAR'], [
        ('DEMURRAGE', 'B > A ise: (B − A) saat × demurrage saatlik oranı', '______ saat × ______', '____________'),
        ('DISPATCH / DESPATCH', 'B < A ise: (A − B) saat × dispatch saatlik oranı', '______ saat × ______', '____________'),
    ], [1900, 3900, 1800, 1760], font=7.2)

    p = d.add_paragraph('4. ÖDEME / FATURA ÖZETİ'); p.style = 'Heading 2'
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(1)
    table(d, ['SONUÇ', 'TARAF', 'TUTAR', 'DAYANAK'], [
        ('☐ Demurrage  ☐ Dispatch  ☐ Süre tam kullanıldı', '☐ Armatör alacağı  ☐ Kiracı alacağı', '____________', 'C/P md. ____ / Time Sheet'),
        ('Hesap tarihi: ____ / ____ / 20____', 'Fatura / Debit Note No.: __________', 'Kur: __________', 'Ekler: NOR / SOF / Time Sheet'),
    ], [2600, 2700, 1800, 2260], font=7.2)

    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    r = p.add_run('Uygulama Notu: '); r.bold = True
    p.add_run('Pozitif zaman farkı demurrage, negatif zaman farkının mutlak değeri dispatch/despatch doğurur. Dispatch yalnızca C/P’de açıkça kararlaştırılmışsa uygulanır. “Once on demurrage, always on demurrage” ve istisna hükümleri ayrıca kontrol edilmelidir.')

    table(d, ['HAZIRLAYAN', 'MASTER / GEMİ ONAYI', 'OPERATION / CHARTERING ONAYI'], [
        ('Ad Soyad / İmza / Tarih:\n\n', 'Ad Soyad / İmza / Tarih:\n\n', 'Ad Soyad / İmza / Tarih:\n\n')
    ], [3000, 3000, 3360], font=7.2)

    path = OUT / 'Sinop_Shipping_Demurrage_Dispatch_Hesaplama_Formu_Siyah_Beyaz.docx'
    d.save(path); print(path)


if __name__ == '__main__':
    build()
