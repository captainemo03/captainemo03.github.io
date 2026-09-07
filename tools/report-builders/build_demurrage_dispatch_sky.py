from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_reports import setup, table, OUT, LOGO_BW, LOGO, NAVY, GRAY

def build():
    d=setup('Demurrage ve Dispatch Hesaplama Formu - SNP SKY','SSC-OPS-DD-01-SKY')
    s=d.sections[0]; s.top_margin=Inches(.46); s.bottom_margin=Inches(.75); s.left_margin=Inches(.65); s.right_margin=Inches(.65)
    d.styles['Normal'].font.size=Pt(8.2); d.styles['Normal'].paragraph_format.space_after=Pt(2)
    logo=LOGO_BW if LOGO_BW.exists() else LOGO
    if logo.exists():
        p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); p.add_run().add_picture(str(logo),width=Inches(.76))
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
    r=p.add_run('DEMURRAGE VE DISPATCH / DESPATCH | SNP SKY'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(14.5); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(4)
    r=p.add_run('KURGUSAL EĞİTİM SENARYOSU • SİYAH-BEYAZ BASKI'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(7.2); r.font.color.rgb=RGBColor.from_string(GRAY)
    table(d,['SEFER / SÖZLEŞME','DEĞER','SEFER / SÖZLEŞME','DEĞER'],[
      ('Gemi','SNP SKY','Sefer / C/P tarihi','SKY-2026/08 | 20.07.2026'),('Yükleme limanı','Hamburg, Almanya','Tahliye limanı','Bilbao, İspanya'),('Yük / Miktar','Çelik rulo / 7.200 MT','Operasyon','Tahliye'),('Laytime şartı','WWD, SHEX EIU','Reversible','Hayır')],[1750,2930,1750,2930],font=7.2)
    p=d.add_paragraph('1. LAYTIME SONUÇ GİRDİLERİ'); p.style='Heading 2'; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(1)
    table(d,['HESAP KALEMİ','GÜN','SAAT','TOPLAM SAAT','AÇIKLAMA'],[
      ('İzin verilen laytime (A)','3 gün','0','72,0','7.200 MT ÷ 2.400 MT/gün'),('Kullanılan net laytime (B)','3 gün','6,5','78,5','Yağmur istisnası düşüldükten sonra'),('Zaman farkı (B − A)','0 gün','+6,5','+6,5','Pozitif fark: Demurrage')],[2650,1000,1000,1600,3110],font=7.4)
    p=d.add_paragraph('2. SÖZLEŞMESEL ORANLAR'); p.style='Heading 2'; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(1)
    table(d,['ORAN','GÜNLÜK TUTAR','SAATLİK TUTAR','PARA BİRİMİ / NOT'],[
      ('Demurrage','8.400 USD/gün','8.400 ÷ 24 = 350 USD/saat','USD'),('Dispatch / Despatch','4.200 USD/gün','4.200 ÷ 24 = 175 USD/saat','USD — bu hesapta uygulanmaz')],[2200,2100,2800,2260],font=7.5)
    p=d.add_paragraph('3. NİHAİ TUTAR HESABI'); p.style='Heading 2'; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(1)
    table(d,['SONUÇ TÜRÜ','UYGULANACAK FORMÜL','HESAP','NİHAİ TUTAR'],[
      ('DEMURRAGE','B > A: (B − A) × saatlik oran','6,5 × 350 USD','2.275 USD'),('DISPATCH / DESPATCH','B < A olması gerekir','78,5 > 72,0','Uygulanmaz')],[1900,3900,1800,1760],font=7.2)
    p=d.add_paragraph('4. ÖDEME / FATURA ÖZETİ'); p.style='Heading 2'; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(1)
    table(d,['SONUÇ','TARAF','TUTAR','DAYANAK'],[
      ('☒ Demurrage  ☐ Dispatch  ☐ Süre tam','☒ Armatör alacağı  ☐ Kiracı alacağı','2.275 USD','C/P + NOR + SOF + Time Sheet'),('Hesap tarihi: 07.08.2026','Debit Note No.: SKY-DM-2026-08','Kur: Uygulanmaz','Ekler kontrol edilmelidir')],[2600,2700,1800,2260],font=7.2)
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2)
    r=p.add_run('Hesap Kontrolü: '); r.bold=True
    p.add_run('İzin verilen 72 saatten net 78,5 saat kullanılmıştır. 6,5 saat aşım × 350 USD/saat = 2.275 USD demurrage. Süre tasarrufu bulunmadığından dispatch/despatch doğmamıştır. Bu kayıt yalnızca eğitim amaçlı kurgusal örnektir; gerçek şirket alacağı değildir.')
    table(d,['HAZIRLAYAN','MASTER / GEMİ ONAYI','OPERATION / CHARTERING ONAYI'],[('Eyüp Miraç OY\nTarih / İmza: __________','Ad Soyad / İmza:\n________________','Ad Soyad / İmza:\n________________')],[3000,3000,3360],font=7.2)
    path=OUT/'SNP_SKY_Demurrage_Dispatch_Doldurulmus_Ornek.docx'; d.save(path); print(path)

if __name__=='__main__': build()
