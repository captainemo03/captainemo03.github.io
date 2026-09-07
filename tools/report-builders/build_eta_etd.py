from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_reports import setup, table, OUT, LOGO_BW, LOGO, NAVY, GRAY


def build():
    d = setup('ETA–ETD Hesaplama Formu', 'SSC-OPS-ETA-01')
    s = d.sections[0]
    s.top_margin = Inches(0.48)
    s.bottom_margin = Inches(0.75)
    s.left_margin = Inches(0.65)
    s.right_margin = Inches(0.65)
    d.styles['Normal'].font.size = Pt(8.5)
    d.styles['Normal'].paragraph_format.space_after = Pt(2)

    logo = LOGO_BW if LOGO_BW.exists() else LOGO
    if logo.exists():
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        p.add_run().add_picture(str(logo), width=Inches(0.82))

    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run('ETA – ETD HESAPLAMA FORMU')
    r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(17); r.font.color.rgb = RGBColor.from_string(NAVY)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('ŞİRKET İÇİ • EĞİTİM AMAÇLI • SİYAH-BEYAZ BASKI')
    r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(7.5); r.font.color.rgb = RGBColor.from_string(GRAY)

    table(d, ['BELGE / SEFER BİLGİSİ', 'DEĞER', 'BELGE / SEFER BİLGİSİ', 'DEĞER'], [
        ('Belge Kodu', 'SSC-OPS-ETA-01', 'Tarih / Revizyon', '____ / ____ / 20____  |  00'),
        ('Gemi', '☐ SNP GALAXY  ☐ SNP SKY  ☐ SNP SPACE', 'Devam', '☐ SNP STAR  ☐ SNP WIND  ☐ SNP PAZAR'),
        ('Kalkış Limanı', '', 'Varış Limanı', ''),
        ('Sefer No.', '', 'Saat Dilimi', 'UTC ____ / LT ____'),
    ], [1900, 2780, 1900, 2780], font=7.6)

    p = d.add_paragraph('1. ETA HESABI — TAHMİNİ VARIŞ ZAMANI')
    p.style = 'Heading 2'; p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    table(d, ['GİRDİ', 'DEĞER', 'BİRİM / AÇIKLAMA'], [
        ('Kalkış tarihi ve saati', '____ / ____ / 20____   ____:____', 'UTC / LT (birini işaretleyin)'),
        ('Toplam rota mesafesi (D)', '', 'Deniz mili — NM'),
        ('Ortalama seyir sürati (V)', '', 'Knot — NM/saat'),
        ('Hava/deniz gecikmesi', '', 'Saat'),
        ('Trafik / kanal / boğaz gecikmesi', '', 'Saat'),
        ('Diğer emniyet payı', '', 'Saat'),
    ], [3200, 2800, 3360], font=7.6)

    table(d, ['HESAPLAMA ADIMI', 'FORMÜL', 'SONUÇ'], [
        ('Saf seyir süresi', 'D ÷ V', '________ saat'),
        ('Toplam gecikme', 'Hava + Trafik + Diğer', '________ saat'),
        ('Düzeltilmiş seyir süresi', '(D ÷ V) + Toplam gecikme', '________ saat'),
        ('ETA', 'Kalkış zamanı + Düzeltilmiş seyir süresi', '____ / ____ / 20____   ____:____'),
    ], [2800, 3900, 2660], font=7.5)

    p = d.add_paragraph('2. ETD HESABI — TAHMİNİ KALKIŞ ZAMANI')
    p.style = 'Heading 2'; p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    table(d, ['LİMAN OPERASYONU', 'TAHMİNİ SÜRE (SAAT)', 'AÇIKLAMA'], [
        ('Varış / yanaşma beklemesi', '', ''),
        ('Yanaşma ve hazırlık', '', ''),
        ('Yükleme / tahliye operasyonu', '', ''),
        ('Survey / evrak / clearance', '', ''),
        ('Kalkış hazırlığı ve emniyet payı', '', ''),
    ], [3300, 2500, 3560], font=7.3)

    table(d, ['HESAPLAMA ADIMI', 'FORMÜL', 'SONUÇ'], [
        ('Toplam liman süresi', 'Tüm liman operasyon sürelerinin toplamı', '________ saat'),
        ('ETD', 'ETA + Toplam liman süresi', '____ / ____ / 20____   ____:____'),
    ], [2800, 3900, 2660], font=7.5)

    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(2)
    r = p.add_run('Kontrol Notu: '); r.bold = True
    p.add_run('Tüm tarih/saat değerlerini aynı saat dilimine çevirin. Sürat veya rota değiştiğinde ETA’yı; liman programı değiştiğinde ETD’yi revize edin. Sonuçlar operasyonel tahmindir ve Master/Operation teyidine tabidir.')

    table(d, ['HAZIRLAYAN', 'MASTER / GEMİ ONAYI', 'OPERATION ONAYI'], [
        ('Ad Soyad / İmza / Tarih:\n\n', 'Ad Soyad / İmza / Tarih:\n\n', 'Ad Soyad / İmza / Tarih:\n\n')
    ], [3120, 3120, 3120], font=7.4)

    path = OUT / 'Sinop_Shipping_ETA_ETD_Hesaplama_Formu_Siyah_Beyaz.docx'
    d.save(path)
    print(path)


if __name__ == '__main__':
    build()
