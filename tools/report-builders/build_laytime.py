from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_reports import setup, table, OUT, LOGO_BW, LOGO, NAVY, GRAY


def build():
    d = setup('Laytime Hesaplama Formu', 'SSC-OPS-LT-01')
    s = d.sections[0]
    s.top_margin = Inches(0.42); s.bottom_margin = Inches(0.72)
    s.left_margin = Inches(0.62); s.right_margin = Inches(0.62)
    d.styles['Normal'].font.size = Pt(8.1)
    d.styles['Normal'].paragraph_format.space_after = Pt(1.5)

    logo = LOGO_BW if LOGO_BW.exists() else LOGO
    if logo.exists():
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(logo), width=Inches(0.72))
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run('LAYTIME HESAPLAMA FORMU')
    r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(16.5); r.font.color.rgb = RGBColor.from_string(NAVY)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run('ŞİRKET İÇİ • EĞİTİM AMAÇLI • SİYAH-BEYAZ BASKI')
    r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(7.2); r.font.color.rgb = RGBColor.from_string(GRAY)

    table(d, ['SEFER / SÖZLEŞME', 'DEĞER', 'SEFER / SÖZLEŞME', 'DEĞER'], [
        ('Gemi', '☐ GALAXY ☐ SKY ☐ SPACE', 'Devam', '☐ STAR ☐ WIND ☐ PAZAR'),
        ('Sefer / C/P tarihi', '', 'Liman / Terminal', ''),
        ('Yük / Miktar', '________________ MT', 'Operasyon', '☐ Yükleme  ☐ Tahliye'),
        ('Laytime şartı', '☐ SHINC  ☐ SHEX  ☐ Diğer: ____', 'Reversible', '☐ Evet  ☐ Hayır'),
    ], [1760, 2920, 1760, 2920], font=7.2)

    p = d.add_paragraph('1. SÖZLEŞMESEL GİRDİLER'); p.style = 'Heading 2'
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    table(d, ['GİRDİ', 'DEĞER', 'HESAP / AÇIKLAMA'], [
        ('Yük miktarı (Q)', '__________ MT', 'Konşimento / C/P miktarı'),
        ('Yükleme-tahliye oranı (R)', '__________ MT/gün', 'C/P’de belirtilen günlük oran'),
        ('İzin verilen laytime', '______ gün  ______ saat', 'Q ÷ R; gün × 24 = toplam saat'),
        ('Demurrage oranı', '__________ USD/gün', 'Saatlik oran = günlük oran ÷ 24'),
        ('Despatch oranı', '__________ USD/gün', 'C/P hükmü; çoğunlukla demurrage’in 1/2’si'),
    ], [3000, 2500, 3860], font=7.2)

    p = d.add_paragraph('2. NOR VE LAYTIME BAŞLANGICI'); p.style = 'Heading 2'
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    table(d, ['OLAY', 'TARİH', 'SAAT', 'NOT'], [
        ('Gemi varışı / All Fast', '', '', ''),
        ('NOR tendered', '', '', '☐ Geçerli  ☐ Şartlı'),
        ('NOR accepted', '', '', ''),
        ('Laytime başlangıcı', '', '', 'C/P notice time uygulandı'),
        ('Operasyon tamamlandı', '', '', 'Loading/Discharging completed'),
    ], [2800, 1700, 1400, 3460], font=7.1)

    p = d.add_paragraph('3. TIME SHEET — SAYILAN / SAYILMAYAN SÜRELER'); p.style = 'Heading 2'
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    table(d, ['BAŞLANGIÇ', 'BİTİŞ', 'OLAY / NEDEN', 'TOPLAM SAAT', 'SAYILAN %', 'LAYTIME SAATİ'], [
        ('____/____  ____:____', '____/____  ____:____', '', '', '100 / 50 / 0', ''),
        ('____/____  ____:____', '____/____  ____:____', '', '', '100 / 50 / 0', ''),
        ('____/____  ____:____', '____/____  ____:____', '', '', '100 / 50 / 0', ''),
        ('____/____  ____:____', '____/____  ____:____', '', '', '100 / 50 / 0', ''),
        ('____/____  ____:____', '____/____  ____:____', '', '', '100 / 50 / 0', ''),
    ], [1500, 1500, 2760, 1200, 1100, 1300], font=6.5)

    p = d.add_paragraph('4. SONUÇ HESABI'); p.style = 'Heading 2'
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    table(d, ['HESAP KALEMİ', 'FORMÜL', 'SONUÇ'], [
        ('Kullanılan net laytime', 'Σ (Toplam saat × Sayılan %)', '__________ saat'),
        ('Süre farkı', 'Kullanılan − İzin verilen', '__________ saat'),
        ('Demurrage', 'Pozitif fark × (Demurrage/gün ÷ 24)', '__________ USD'),
        ('Despatch', 'Negatif farkın mutlak değeri × (Despatch/gün ÷ 24)', '__________ USD'),
    ], [2600, 4400, 2360], font=7.0)

    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(1)
    r = p.add_run('Kontrol: '); r.bold = True
    p.add_run('C/P’nin NOR, notice time, WIBON/WIPON, SHINC/SHEX, weather working day, shifting, strike ve “once on demurrage, always on demurrage” hükümleri ayrıca kontrol edilmelidir. Bu form C/P yorumunun yerine geçmez.')

    table(d, ['HAZIRLAYAN', 'MASTER / GEMİ ONAYI', 'OPERATION ONAYI'], [
        ('Ad Soyad / İmza / Tarih:\n', 'Ad Soyad / İmza / Tarih:\n', 'Ad Soyad / İmza / Tarih:\n')
    ], [3120, 3120, 3120], font=7.0)

    path = OUT / 'Sinop_Shipping_Laytime_Hesaplama_Formu_Siyah_Beyaz.docx'
    d.save(path); print(path)


if __name__ == '__main__':
    build()
