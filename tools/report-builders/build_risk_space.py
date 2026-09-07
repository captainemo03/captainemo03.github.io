from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_reports import setup, table, OUT, LOGO_BW, LOGO, NAVY, GRAY

def build():
    d=setup('Sefer Bazlı Risk Analizi - SNP SPACE','SSC-RA-SPC-2026-01')
    s=d.sections[0]; s.top_margin=Inches(.42); s.bottom_margin=Inches(.74); s.left_margin=Inches(.55); s.right_margin=Inches(.55)
    d.styles['Normal'].font.size=Pt(7.8); d.styles['Normal'].paragraph_format.space_after=Pt(1.5)
    logo=LOGO_BW if LOGO_BW.exists() else LOGO
    if logo.exists():
        p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); p.add_run().add_picture(str(logo),width=Inches(.68))
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
    r=p.add_run('SEFER BAZLI RİSK ANALİZİ | SNP SPACE'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(15.5); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(3)
    r=p.add_run('KURGUSAL EĞİTİM SENARYOSU • SİYAH-BEYAZ BASKI'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(7.1); r.font.color.rgb=RGBColor.from_string(GRAY)
    table(d,['SEFER BİLGİSİ','DEĞER','SEFER BİLGİSİ','DEĞER'],[
      ('Gemi','SNP SPACE','Belge / Revizyon','SSC-RA-SPC-2026-01 / 00'),('Rota','Antwerp, Belçika → Mersin, Türkiye','Sefer tarihi','10–23 Ağustos 2026'),('Yük','1.850 MT proje yükü / çelik konstrüksiyon','Kritik geçişler','İngiliz Kanalı, Cebelitarık, Akdeniz'),('Hava varsayımı','Kuzey Denizi: 25–30 kn rüzgâr','Değerlendirme','Sefer öncesi / operasyon öncesi')],[1700,3180,1700,2780],font=6.9)
    p=d.add_paragraph('PUANLAMA: Olasılık (O) 1–5 × Etki (E) 1–5 = Risk. 1–4 Düşük | 5–9 Orta | 10–16 Yüksek | 17–25 Kritik')
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].bold=True; p.paragraph_format.space_after=Pt(2)
    table(d,['NO','TEHLİKE / RİSK','MEVCUT DURUM / NEDEN','O','E','İLK RİSK','EK KONTROLLER','KALAN RİSK'],[
      ('1','Proje yükünün kayması veya bağların çözülmesi','Ağır ve yüksek ağırlık merkezli parçalar; Kuzey Denizi hava riski','4','5','20 Kritik','Onaylı lashing planı; kalkış öncesi bağımsız kontrol; ilk 12 saatte ve ağır hava öncesi tekrar kontrol','2×5=10 Yüksek'),
      ('2','Çatışma / yakın geçiş','İngiliz Kanalı ve Cebelitarık’ta yoğun trafik','3','5','15 Yüksek','Master passage plan review; ek gözcü; radar/ECDIS/visual çapraz kontrol; CPA/TCPA limitleri','2×5=10 Yüksek'),
      ('3','Ağır hava ve güverteye deniz basması','25–30 kn rüzgâr; proje yükünün rüzgâr alanı yüksek','3','5','15 Yüksek','Weather routing; emniyetli sürat/rota; kapak ve freeing port kontrolleri; heavy-weather checklist','2×4=8 Orta'),
      ('4','Ana makine veya dümen arızası','Uzun sefer ve kritik dar su geçişleri','2','5','10 Yüksek','Kritik ekipman testi; yedek parça kontrolü; geçiş öncesi engine/steering test ve standby','1×5=5 Orta'),
      ('5','Yorgunluk / vardiya hatası','Yoğun trafik bölgelerinde uzun köprüüstü iş yükü','4','4','16 Yüksek','Work-rest ön kontrolü; Master çağırma kriterleri; kritik geçişte vardiya takviyesi','2×4=8 Orta'),
      ('6','Yük operasyonunda ezilme / düşen cisim','Vinç operasyonu, askı altında çalışma ve kör noktalar','3','5','15 Yüksek','Lift plan; exclusion zone; banksman; toolbox; sertifikalı sapan ve iletişim testi','1×5=5 Orta'),
      ('7','Bunker sırasında yakıt döküntüsü','Transfer bağlantısı, tank overfill veya iletişim hatası','2','4','8 Orta','SOPEP hazır; scupper kapalı; drip tray; tank sounding; acil stop testi ve sürekli haberleşme','1×4=4 Düşük'),
      ('8','GNSS/AIS sapması veya siber kesinti','Elektronik seyir girdilerine aşırı bağımlılık','3','4','12 Yüksek','Radar/visual fix; bağımsız pozisyon doğrulama; erişim kontrolü; yedekleme ve olay raporu','2×3=6 Orta')],[420,1350,2000,300,300,760,3050,1180],font=5.75)
    p=d.add_paragraph('ÖNCELİKLİ AKSİYONLAR'); p.style='Heading 2'; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(1)
    table(d,['AKSİYON','SORUMLU','TERMİN / KANIT'],[
      ('Lashing planı ve ağır hava limitlerini yükleme bitmeden doğrula','Master / Chief Officer','Kalkış öncesi • İmzalı checklist ve fotoğraf'),('Kritik geçiş planını köprüüstü ekibiyle brifing et','Master','Her kritik geçiş öncesi • BRM kayıt formu'),('Kritik makine/dümen testleri ile yedek durumunu kaydet','Chief Engineer','Kalkış ve Cebelitarık öncesi • Log/PMS kaydı')],[5200,1800,2360],font=6.6)
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(1)
    r=p.add_run('Sonuç: '); r.bold=True
    p.add_run('Ek kontroller uygulandıktan sonra kritik risk kalmamaktadır. Yük kayması ve çatışma riskleri 10 puanla “Yüksek” seviyede kaldığından Master gözetimi ve sefer boyunca dinamik yeniden değerlendirme gerektirir. Bu çalışma gerçek SNP SPACE seferi veya şirket risk kaydı değildir.')
    table(d,['HAZIRLAYAN','MASTER / GEMİ ONAYI','OPERATION / TECHNICAL ONAYI'],[('Eyüp Miraç OY\nTarih / İmza: __________','Ad Soyad / İmza:\n________________','Ad Soyad / İmza:\n________________')],[3000,3000,3360],font=6.8)
    path=OUT/'SNP_SPACE_Risk_Analizi_Doldurulmus_Ornek.docx'; d.save(path); print(path)

if __name__=='__main__': build()
